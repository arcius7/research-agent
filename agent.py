"""
Research RAG Agent  —  LangGraph multi-node pipeline
=====================================================

Core idea
─────────
When a paper is ingested the agent analyses its size and derives three things
automatically:

  1. Chunk size   — finer chunks for long papers (better retrieval precision),
                    coarser chunks for short papers (better context coverage).

  2. Pomodoro     — work-session duration set to ≈ 2 min/page so the timer
     duration       scales with how long the paper actually takes to read.

  3. Voice        — a speaker ID is picked deterministically from the paper's
                    filename, giving each paper a consistent, distinct voice
                    (rendered by macOS `say` in tts.py).

Nodes
─────
  ingest     PDF / Word / Excel / JSON → adaptive chunks → nomic-embed-text
             → turbovec 4-bit quantized store → updates Pomodoro timer

  retrieve   ANN similarity search → fills context

  llm        RAG answer over retrieved context (local Ollama, or a cloud
             provider when its API key is set in .env)

  pomodoro   Pomodoro timer (25/5/15 min default, overridden by paper size)

Graph routing
─────────────
  START → dispatch
    "ingest"        → ingest → END
    "query"         → retrieve → llm → END
    "search"        → retrieve → END
    timer actions   → pomodoro → END
"""

from __future__ import annotations

import hashlib
import json
import os
import threading
import time
from pathlib import Path
from typing import Optional, TypedDict

from dotenv import load_dotenv
load_dotenv(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env"))

import requests
from langgraph.graph import StateGraph, START, END

from logging_setup import get_logger
log = get_logger(__name__)

_HERE = os.path.dirname(os.path.abspath(__file__))

# ── turbovec + LangChain (LCEL) ───────────────────────────────────────────────
from turbovec.langchain import TurboQuantVectorStore  # noqa: E402
from langchain_core.prompts import PromptTemplate       # noqa: E402
from langchain_core.runnables import RunnableLambda, RunnablePassthrough  # noqa: E402

# ── shared pomodoro state (dedicated module — NOT server.py, which would load
#    twice and split the state when the server runs as __main__) ───────────────
from timer_state import (  # noqa: E402
    state as _timer, state_lock,
    advance_mode as _advance_mode, MODE_KEY as _MODE_KEY,
)

# ═══════════════════════════════════════════════════════════════════════════════
# Static config
# ═══════════════════════════════════════════════════════════════════════════════

OLLAMA_BASE  = os.environ.get("OLLAMA_BASE", "http://localhost:11434")
LLM_MODEL    = os.environ.get("LLM_MODEL", "gemma4:e4b")
EMBED_MODEL  = os.environ.get("EMBED_MODEL", "nomic-embed-text")

TURBOVEC_DIR = os.path.join(_HERE, ".turbovec_store")

# Single-paper mode (default ON): the agent holds exactly ONE paper at a time.
# Uploading a paper wipes the store so questions answer against only that paper —
# no bleed from previously uploaded papers. Set SINGLE_PAPER=0 for a shared store.
SINGLE_PAPER = os.environ.get("SINGLE_PAPER", "1") not in ("0", "false", "False")

# ── Cloud LLM provider detection ─────────────────────────────────────────────
# When an API key is set in .env, we route LLM calls through that provider
# instead of local Ollama. Embeddings always stay on Ollama (local, free).

_PROVIDERS: dict[str, dict] = {}

def _detect_providers() -> None:
    """Scan env for known API keys and register available cloud providers."""
    _PROVIDERS.clear()
    if os.environ.get("OPENAI_API_KEY"):
        _PROVIDERS["openai"] = {
            "key": os.environ["OPENAI_API_KEY"],
            "url": "https://api.openai.com/v1/chat/completions",
            "models": ["gpt-4o", "gpt-4o-mini", "gpt-4.1", "gpt-4.1-mini", "gpt-4.1-nano", "o3-mini"],
        }
    if os.environ.get("ANTHROPIC_API_KEY"):
        _PROVIDERS["anthropic"] = {
            "key": os.environ["ANTHROPIC_API_KEY"],
            "url": "https://api.anthropic.com/v1/messages",
            "models": ["claude-sonnet-4-6", "claude-haiku-4-5-20251001"],
        }
    if os.environ.get("GEMINI_API_KEY"):
        _PROVIDERS["gemini"] = {
            "key": os.environ["GEMINI_API_KEY"],
            "url": "https://generativelanguage.googleapis.com/v1beta/chat/completions",
            "models": ["gemini-2.5-flash", "gemini-2.5-pro"],
        }
    if os.environ.get("GROQ_API_KEY"):
        _PROVIDERS["groq"] = {
            "key": os.environ["GROQ_API_KEY"],
            "url": "https://api.groq.com/openai/v1/chat/completions",
            "models": ["llama-3.3-70b-versatile", "llama-3.1-8b-instant", "mixtral-8x7b-32768"],
        }
    if os.environ.get("TOGETHER_API_KEY"):
        _PROVIDERS["together"] = {
            "key": os.environ["TOGETHER_API_KEY"],
            "url": "https://api.together.xyz/v1/chat/completions",
            "models": ["meta-llama/Llama-3.3-70B-Instruct-Turbo", "mistralai/Mixtral-8x7B-Instruct-v0.1"],
        }

_detect_providers()


def _provider_for_model(model: str) -> tuple[Optional[str], Optional[dict]]:
    """Return (provider_name, provider_info) if model belongs to a cloud provider."""
    for name, info in _PROVIDERS.items():
        if model in info["models"]:
            return name, info
    return None, None


# Cloud models aren't thermally constrained — allow much longer answers than
# the local 512-token cap (LLM_NUM_PRED below).
CLOUD_NUM_PRED = int(os.environ.get("CLOUD_NUM_PRED", "2048"))


def _cloud_headers(provider: str, info: dict) -> dict:
    if provider == "anthropic":
        return {"x-api-key": info["key"],
                "anthropic-version": "2023-06-01",
                "content-type": "application/json"}
    # openai, gemini, groq, together — all OpenAI-compatible
    return {"Authorization": f"Bearer {info['key']}",
            "Content-Type": "application/json"}


def _cloud_generate(prompt: str, model: str, provider: str, info: dict,
                    max_tokens: Optional[int] = None) -> str:
    """Call a cloud LLM's chat completions API (Anthropic or OpenAI-compatible)."""
    resp = requests.post(
        info["url"],
        headers=_cloud_headers(provider, info),
        json={
            "model": model,
            "messages": [{"role": "user", "content": prompt}],
            "max_tokens": max_tokens or CLOUD_NUM_PRED,
        },
        timeout=120,
    )
    resp.raise_for_status()
    data = resp.json()
    if provider == "anthropic":
        return data["content"][0]["text"]
    return data["choices"][0]["message"]["content"]


def _cloud_stream(prompt: str, model: str, provider: str, info: dict):
    """Stream tokens from a cloud LLM. Yields text chunks."""
    resp = requests.post(
        info["url"],
        headers=_cloud_headers(provider, info),
        json={
            "model": model,
            "messages": [{"role": "user", "content": prompt}],
            "max_tokens": CLOUD_NUM_PRED,
            "stream": True,
        },
        timeout=300,
        stream=True,
    )
    resp.raise_for_status()
    for line in resp.iter_lines():
        if not line:
            continue
        text = line.decode()
        if not text.startswith("data: "):
            continue
        data = text[6:]
        if data.strip() == "[DONE]":
            break
        try:
            obj = json.loads(data)
        except json.JSONDecodeError:
            continue
        if provider == "anthropic":
            if obj.get("type") == "content_block_delta":
                tok = obj.get("delta", {}).get("text", "")
            else:
                tok = ""
        else:
            tok = obj.get("choices", [{}])[0].get("delta", {}).get("content", "") or ""
        if tok:
            yield tok

# Voice pool size — speaker_id is an index into tts.py's curated `say` voices.
_N_VOICES = 21


# ═══════════════════════════════════════════════════════════════════════════════
# Paper profile  —  the heart of the adaptive logic
# ═══════════════════════════════════════════════════════════════════════════════

def paper_profile(filename: str, total_chars: int, n_pages: int) -> dict:
    """
    Derive chunk_size, overlap, Pomodoro work duration, and voice speaker ID
    from a paper's size metrics.

    Returns
    -------
    dict with keys:
      chunk_size    int   characters per embedding chunk
      overlap       int   overlap between consecutive chunks
      work_minutes  int   Pomodoro work-session length in minutes
      work_seconds  int   same in seconds (used to set the timer)
      speaker_id    int   voice index (0-20), rendered by tts.py
      total_chars   int
      n_pages       int
    """
    # ── chunk size: bigger chunks = far fewer embeddings (lighter on 16 GB) ────
    # Still scales down a bit for longer papers, but stays coarse to keep the
    # number of vectors (and embedding calls) low.
    if total_chars < 8_000:        # ~4 pages
        chunk_size, overlap = 1800, 200
    elif total_chars < 25_000:     # ~15 pages
        chunk_size, overlap = 1500, 180
    elif total_chars < 60_000:     # ~35 pages
        chunk_size, overlap = 1200, 150
    elif total_chars < 120_000:    # ~70 pages
        chunk_size, overlap = 1000, 120
    else:                           # thesis / book chapter
        chunk_size, overlap = 850, 100

    # ── Pomodoro: 2 min/page, clamped to [15, 50] min ────────────────────────
    work_minutes = max(15, min(50, n_pages * 2))

    # ── voice: deterministic hash of filename → consistent per paper ──────────
    h = int(hashlib.md5(filename.encode()).hexdigest()[:8], 16)
    speaker_id = h % _N_VOICES

    return {
        "chunk_size":   chunk_size,
        "overlap":      overlap,
        "work_minutes": work_minutes,
        "work_seconds": work_minutes * 60,
        "speaker_id":   speaker_id,
        "total_chars":  total_chars,
        "n_pages":      n_pages,
    }


# ═══════════════════════════════════════════════════════════════════════════════
# Graph state
# ═══════════════════════════════════════════════════════════════════════════════

class AgentState(TypedDict):
    # timer
    action:        str
    timer:         dict
    announce:      Optional[str]   # timer-transition text (spoken by tts.py)
    audio_path:    Optional[str]

    # RAG
    files:         list            # paths for "ingest"
    query:         Optional[str]   # question for "query" / "search"
    paper:         Optional[str]   # restrict retrieval to this paper (source filter)
    context:       list            # retrieved chunks
    answer:        Optional[str]   # LLM response
    results:       list            # raw search hits

    # paper profile (set by ingest_node, carried through the run)
    speaker_id:    Optional[int]   # voice index for this paper (tts.py)
    paper_profile: Optional[dict]  # full profile dict


# ═══════════════════════════════════════════════════════════════════════════════
# Ollama helpers
# ═══════════════════════════════════════════════════════════════════════════════

# ── Lightweight tuning for 16 GB Macs ─────────────────────────────────────────
EMBED_BATCH  = 24      # embed this many chunks per request (steady memory)
LLM_NUM_PRED = 512     # cap local answer length (cloud uses CLOUD_NUM_PRED)
LLM_KEEPALIVE = "5m"   # unload the model after idle to free RAM


def _ollama_embed(texts: list[str]) -> list[list[float]]:
    """Embed in small batches so a big paper doesn't spike memory."""
    out: list[list[float]] = []
    for i in range(0, len(texts), EMBED_BATCH):
        resp = requests.post(
            f"{OLLAMA_BASE}/api/embed",
            json={"model": EMBED_MODEL, "input": texts[i:i + EMBED_BATCH]},
            timeout=120,
        )
        resp.raise_for_status()
        out.extend(resp.json()["embeddings"])
    return out


def set_model(name: str) -> None:
    """Switch the LLM used for answers at runtime (from the UI dropdown)."""
    global LLM_MODEL
    if name:
        LLM_MODEL = name


def get_model() -> str:
    return LLM_MODEL


def list_models() -> dict:
    """All available models, grouped: {"local": [...], "cloud": [...]}.
    Re-detects providers so a freshly edited .env only needs a page refresh."""
    load_dotenv(os.path.join(_HERE, ".env"), override=True)
    _detect_providers()
    cloud = [m for info in _PROVIDERS.values() for m in info["models"]]
    local = []
    try:
        r = requests.get(f"{OLLAMA_BASE}/api/tags", timeout=10)
        r.raise_for_status()
        local = [m["name"] for m in r.json().get("models", [])]
    except requests.RequestException:
        pass
    return {"local": local, "cloud": cloud}


def generate(prompt: str, max_tokens: Optional[int] = None) -> str:
    """Generate a completion with the active model. Routes to a cloud provider
    if the active model belongs to one, else local Ollama (with retry-on-empty,
    since Ollama occasionally returns nothing on the call that reloads a model)."""
    provider, info = _provider_for_model(LLM_MODEL)
    _t0 = time.time()
    if provider and info:
        log.info("generate via CLOUD %s/%s", provider, LLM_MODEL)
        out = _cloud_generate(prompt, LLM_MODEL, provider, info, max_tokens)
        log.info("generate done (%s) %d chars in %.1fs", provider, len(out), time.time() - _t0)
        return out

    log.info("generate via OLLAMA %s", LLM_MODEL)
    payload = {
        "model": LLM_MODEL,
        "messages": [{"role": "user", "content": prompt}],
        "stream": False,
        "keep_alive": LLM_KEEPALIVE,
        "options": {"num_predict": max_tokens or LLM_NUM_PRED},
    }
    text = ""
    for attempt in range(2):
        resp = requests.post(f"{OLLAMA_BASE}/api/chat", json=payload, timeout=300)
        resp.raise_for_status()
        text = resp.json().get("message", {}).get("content", "")
        if text.strip():
            break
        log.warning("empty Ollama response (attempt %d) — retrying", attempt + 1)
    log.info("generate done (ollama) %d chars in %.1fs", len(text), time.time() - _t0)
    return text


_generate = generate  # internal alias used by the RAG chain


class _OllamaEmbeddings:
    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return _ollama_embed(texts)

    def embed_query(self, text: str) -> list[float]:
        return _ollama_embed([text])[0]


# ═══════════════════════════════════════════════════════════════════════════════
# Node: pomodoro
# ═══════════════════════════════════════════════════════════════════════════════

_TRANSITION_TEXT = {
    "short_break": "Nice work! Take a short break.",
    "long_break":  "Great effort! Enjoy a long break.",
    "work":        "Break over. Time to focus.",
}


def pomodoro_node(state: AgentState) -> AgentState:
    action = state.get("action", "tick")
    announce: Optional[str] = None

    with state_lock:
        if action == "start":
            _timer["running"] = True
        elif action == "pause":
            _timer["running"] = False
        elif action == "reset":
            _timer["running"] = False
            _timer["elapsed"] = 0
        elif action == "skip":
            _advance_mode()
            announce = _TRANSITION_TEXT[_timer["mode"]]
        elif action == "tick" and _timer["running"]:
            _timer["elapsed"] += 1
            duration = _timer["settings"][_MODE_KEY[_timer["mode"]]]
            if _timer["elapsed"] >= duration:
                _advance_mode()
                announce = _TRANSITION_TEXT[_timer["mode"]]
        snapshot = {**_timer, "tasks": list(_timer["tasks"])}

    return {**state, "timer": snapshot, "announce": announce}


# ═══════════════════════════════════════════════════════════════════════════════
# Node: ingest
# ═══════════════════════════════════════════════════════════════════════════════

# ── file extractors ───────────────────────────────────────────────────────────

def _extract_pdf(path: str) -> list[tuple[str, dict]]:
    import pypdf
    fname  = Path(path).name
    reader = pypdf.PdfReader(path)
    pages  = [
        (page.extract_text(), {"source": fname, "page": i})
        for i, page in enumerate(reader.pages)
        if page.extract_text()
    ]
    return pages


def _extract_docx(path: str) -> list[tuple[str, dict]]:
    import docx
    fname = Path(path).name
    doc   = docx.Document(path)
    return [(p.text, {"source": fname}) for p in doc.paragraphs if p.text.strip()]


def _extract_excel(path: str) -> list[tuple[str, dict]]:
    import openpyxl
    fname = Path(path).name
    wb    = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out   = []
    for sheet in wb.worksheets:
        rows = [
            " | ".join(str(c) for c in row if c is not None)
            for row in sheet.iter_rows(values_only=True)
        ]
        rows = [r for r in rows if r.strip()]
        if rows:
            out.append(("\n".join(rows), {"source": fname, "sheet": sheet.title}))
    return out


def _extract_json(path: str) -> list[tuple[str, dict]]:
    fname = Path(path).name
    with open(path) as f:
        data = json.load(f)
    items = data if isinstance(data, list) else [data]
    return [(json.dumps(item, ensure_ascii=False), {"source": fname}) for item in items]


_EXTRACTORS = {
    ".pdf":  _extract_pdf,
    ".docx": _extract_docx,
    ".doc":  _extract_docx,
    ".xlsx": _extract_excel,
    ".xls":  _extract_excel,
    ".json": _extract_json,
}

# ── adaptive chunker ──────────────────────────────────────────────────────────

def _chunk(text: str, meta: dict, chunk_size: int, overlap: int) -> list[tuple[str, dict]]:
    """Split text into overlapping chunks sized for this paper."""
    separators = ["\n\n", "\n", ". ", " "]
    chunks, start = [], 0
    n = len(text)
    while start < n:
        end = min(start + chunk_size, n)
        if end < n:
            for sep in separators:
                pos = text.rfind(sep, start + overlap, end)
                if pos != -1:
                    end = pos + len(sep)
                    break
        piece = text[start:end].strip()
        if piece:
            chunks.append((piece, {**meta, "chunk": len(chunks)}))
        if end >= n:                          # reached the end → stop
            break
        start = max(end - overlap, start + 1)  # always make forward progress
    return chunks

# ── vector store ──────────────────────────────────────────────────────────────

_store: Optional[TurboQuantVectorStore] = None
_store_lock = threading.Lock()


def _get_store() -> TurboQuantVectorStore:
    global _store
    if _store is None:
        with _store_lock:
            if _store is None:
                emb  = _OllamaEmbeddings()
                path = Path(TURBOVEC_DIR)
                _store = (
                    TurboQuantVectorStore.load(path, emb)
                    if (path / "index.tvim").exists()
                    else TurboQuantVectorStore(embedding=emb, bit_width=4)
                )
    return _store


def reset_store() -> None:
    """Wipe the vector store (in-memory + on-disk). Single-paper mode calls this
    before ingesting so the store only ever holds the newest paper."""
    global _store
    import shutil
    with _store_lock:
        _store = TurboQuantVectorStore(embedding=_OllamaEmbeddings(), bit_width=4)
        shutil.rmtree(TURBOVEC_DIR, ignore_errors=True)
    log.info("store reset (single-paper mode) — previous embeddings wiped")

# ── ingest node ───────────────────────────────────────────────────────────────

def ingest_node(state: AgentState) -> AgentState:
    # Single-paper mode: clear everything first so this upload becomes the ONLY
    # paper in memory — questions can't bleed in from earlier papers.
    if SINGLE_PAPER:
        reset_store()
    store    = _get_store()
    files    = state.get("files", [])
    ingested = []
    profile  = None

    for path in files:
        ext     = Path(path).suffix.lower()
        extract = _EXTRACTORS.get(ext)
        if extract is None:
            raise ValueError(f"Unsupported file type: {ext}  ({path})")

        raw         = extract(path)
        total_chars = sum(len(t) for t, _ in raw)
        n_pages     = (
            len(raw) if ext == ".pdf"            # one entry per page
            else max(1, total_chars // 2_500)    # estimate for other formats
        )
        fname   = Path(path).name
        profile = paper_profile(fname, total_chars, n_pages)

        # ── chunk with sizes tuned to this paper ──────────────────────────────
        texts, metas, ids = [], [], []
        for text, meta in raw:
            for chunk, cmeta in _chunk(text, meta,
                                       profile["chunk_size"], profile["overlap"]):
                texts.append(chunk)
                metas.append(cmeta)
                # Stable id per (paper, position) → re-uploading the same paper
                # UPSERTS instead of adding a duplicate copy.
                ids.append(f"{fname}::{len(ids)}")

        if texts:
            # Drop this paper's previous chunks first: stable ids upsert the
            # overlap, but a re-upload with FEWER chunks would otherwise leave
            # stale tail chunks polluting retrieval.
            old_ids = [k for k in getattr(store, "_docs", {}) if k.startswith(f"{fname}::")]
            if old_ids:
                try:
                    store.delete(old_ids)
                except Exception:
                    pass
            log.info("embedding %s: %d chunks (%d chars, %d pages) via %s",
                     fname, len(texts), total_chars, n_pages, EMBED_MODEL)
            _t0 = time.time()
            store.add_texts(texts, metadatas=metas, ids=ids)
            log.info("embedded %s: %d chunks in %.1fs",
                     fname, len(texts), time.time() - _t0)
            ingested.append(fname)

        # ── set Pomodoro work duration from paper size ─────────────────────────
        with state_lock:
            _timer["settings"]["work_duration"] = profile["work_seconds"]
            _timer["elapsed"]  = 0
            _timer["running"]  = False
            _timer["mode"]     = "work"

    store.dump(TURBOVEC_DIR)

    # Audio is generated on demand via macOS `say` (tts.py, /api/speak).
    with state_lock:
        snapshot = {**_timer, "tasks": list(_timer["tasks"])}

    return {
        **state,
        "timer":         snapshot,
        "paper_profile": profile,
        "speaker_id":    profile["speaker_id"] if profile else state.get("speaker_id"),
        "announce":      None,
        "results":       [{"ingested": ingested,
                           "total_docs": len(getattr(store, "_docs", {})),
                           "profile": profile}],
    }


# ═══════════════════════════════════════════════════════════════════════════════
# Node: retrieve
# ═══════════════════════════════════════════════════════════════════════════════

def _retrieve(query: str, paper: Optional[str] = None, k: int = 4):
    """Top-k chunks. In single-paper mode the store holds one paper, so we skip
    the source filter entirely (a stale current_paper can't starve retrieval).
    In multi-paper mode, `paper` restricts to that paper's chunks."""
    store   = _get_store()
    flt     = None if SINGLE_PAPER else ({"source": paper} if paper else None)
    hits    = store.similarity_search_with_score(query, k=k, filter=flt)
    top     = round(hits[0][1], 3) if hits else None
    log.info("retrieve: %d hits (top score %s) for %r", len(hits), top, query[:60])
    return hits


def retrieve_node(state: AgentState) -> AgentState:
    hits = _retrieve(state.get("query") or "", state.get("paper"))
    return {
        **state,
        "context": [doc.page_content for doc, _ in hits],
        "results": [
            {"text": doc.page_content, "score": round(score, 4), "meta": doc.metadata}
            for doc, score in hits
        ],
    }


# ═══════════════════════════════════════════════════════════════════════════════
# Node: llm  (Ollama RAG)
# ═══════════════════════════════════════════════════════════════════════════════

_RAG_PROMPT = """\
You are a research assistant that answers questions about academic papers.

Retrieved context from papers:
---
{context}
---

Question: {question}

Instructions:
- Answer using only the provided context.
- Be precise; cite the source document name when relevant.
- If the context is insufficient, say so explicitly.

Answer:"""


def llm_node(state: AgentState) -> AgentState:
    question = state.get("query") or ""
    context  = state.get("context") or []
    if not context:
        return {**state, "answer": "No relevant context found in the document store."}
    prompt = _RAG_PROMPT.format(
        context="\n\n---\n\n".join(context),
        question=question,
    )
    return {**state, "answer": _generate(prompt)}


# ── Idiomatic LCEL RAG chain (retriever → prompt → Ollama) ────────────────────
# Deterministic Runnable pipeline — no agent loop, so it stays fast and cool.

_RAG_TEMPLATE = PromptTemplate.from_template(_RAG_PROMPT)   # {context} {question}


# Cap the retrieved context so prompt-processing stays fast (time-to-first-token).
# Dense scientific text tokenizes heavily; ~2800 chars keeps prompt eval ~4-5 s.
CONTEXT_CHAR_CAP = 2800


def _join_capped(texts: list[str]) -> str:
    """Join chunks up to the context cap (whole chunks only)."""
    out, total = [], 0
    for t in texts:
        if out and total + len(t) > CONTEXT_CHAR_CAP:
            break
        out.append(t)
        total += len(t)
    return "\n\n---\n\n".join(out)


def _format_docs(docs) -> str:
    return _join_capped([d.page_content for d in docs])


def build_rag_chain(paper: Optional[str] = None):
    """LangChain Expression Language chain: retrieve → prompt → LLM. In
    single-paper mode the filter is dropped (store holds one paper)."""
    flt       = None if SINGLE_PAPER else ({"source": paper} if paper else None)
    retriever = _get_store().as_retriever(search_kwargs={"k": 4, "filter": flt})
    return (
        {"context": retriever | RunnableLambda(_format_docs),
         "question": RunnablePassthrough()}
        | _RAG_TEMPLATE
        | RunnableLambda(lambda pv: _generate(pv.to_string()))
    )


def stream_answer(question: str, paper: Optional[str] = None):
    """Generator yielding answer tokens as they arrive (for SSE).
    Routes through cloud provider when the active model is a cloud model."""
    hits    = _retrieve(question, paper)
    context = _join_capped([doc.page_content for doc, _ in hits])
    if not context:
        yield "No relevant context found for this paper."
        return
    prompt = _RAG_PROMPT.format(context=context, question=question)

    provider, info = _provider_for_model(LLM_MODEL)
    if provider and info:
        yield from _cloud_stream(prompt, LLM_MODEL, provider, info)
        return

    with requests.post(
        f"{OLLAMA_BASE}/api/chat",
        json={
            "model": LLM_MODEL,
            "messages": [{"role": "user", "content": prompt}],
            "stream": True,
            "keep_alive": LLM_KEEPALIVE,
            "options": {"num_predict": LLM_NUM_PRED},
        },
        timeout=300,
        stream=True,
    ) as resp:
        resp.raise_for_status()
        for line in resp.iter_lines():
            if not line:
                continue
            obj = json.loads(line)
            tok = obj.get("message", {}).get("content", "")
            if tok:
                yield tok
            if obj.get("done"):
                break


# ═══════════════════════════════════════════════════════════════════════════════
# Routing
# ═══════════════════════════════════════════════════════════════════════════════

_TIMER_ACTIONS = {"start", "pause", "reset", "skip", "tick"}


def _dispatch(state: AgentState) -> str:
    action = state.get("action", "tick")
    if action in _TIMER_ACTIONS:        return "pomodoro"
    if action == "ingest":              return "ingest"
    if action in ("query", "search"):   return "retrieve"
    return END


def _after_retrieve(state: AgentState) -> str:
    return "llm" if state.get("action") == "query" else END


# ═══════════════════════════════════════════════════════════════════════════════
# Graph
# ═══════════════════════════════════════════════════════════════════════════════

def build_graph():
    g = StateGraph(AgentState)

    g.add_node("pomodoro", pomodoro_node)
    g.add_node("ingest",   ingest_node)
    g.add_node("retrieve", retrieve_node)
    g.add_node("llm",      llm_node)

    g.add_conditional_edges(START, _dispatch, {
        "pomodoro": "pomodoro",
        "ingest":   "ingest",
        "retrieve": "retrieve",
        END:        END,
    })
    g.add_conditional_edges("retrieve", _after_retrieve, {"llm": "llm", END: END})

    g.add_edge("pomodoro", END)
    g.add_edge("ingest",   END)
    g.add_edge("llm",      END)

    return g.compile()


agent = build_graph()


# ═══════════════════════════════════════════════════════════════════════════════
# Convenience API
# ═══════════════════════════════════════════════════════════════════════════════

def _base(**kw) -> AgentState:
    return {
        "action": "tick", "timer": {}, "announce": None, "audio_path": None,
        "files": [], "query": None, "paper": None,
        "context": [], "answer": None, "results": [],
        "speaker_id": None, "paper_profile": None,
        **kw,
    }


def quick_profile(path: str) -> dict:
    """FAST path: estimate the profile and set the Pomodoro timer — NO embedding
    and, for PDFs, NO full text extraction (a 300-page PDF would block the
    upload response for seconds). Samples a few pages and extrapolates."""
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        import pypdf
        reader  = pypdf.PdfReader(path)
        n_pages = len(reader.pages)
        # Sample up to 6 pages spread through the document → chars/page estimate
        idx     = sorted({0, n_pages // 4, n_pages // 2, (3 * n_pages) // 4,
                          n_pages - 1, min(1, n_pages - 1)})
        sampled = [len(reader.pages[i].extract_text() or "") for i in idx]
        per_pg  = (sum(sampled) / len(sampled)) if sampled else 2_500
        total_chars = int(per_pg * n_pages)
    else:
        extract = _EXTRACTORS.get(ext)
        if extract is None:
            raise ValueError(f"Unsupported file type: {ext}")
        raw         = extract(path)
        total_chars = sum(len(t) for t, _ in raw)
        n_pages     = max(1, total_chars // 2_500)
    profile = paper_profile(Path(path).name, total_chars, n_pages)
    with state_lock:
        _timer["settings"]["work_duration"] = profile["work_seconds"]
        _timer["elapsed"] = 0
        _timer["running"] = False
        _timer["mode"]    = "work"
    return profile


def ingest(files: list[str]) -> dict:
    """Ingest files (heavy: embeds). Returns result including the paper profile."""
    return agent.invoke(_base(action="ingest", files=files))


def query(question: str, paper: Optional[str] = None) -> str:
    """Ask a question; answer grounded in `paper` (or all papers if None).
    Runs the idiomatic LCEL RAG chain."""
    return build_rag_chain(paper).invoke(question)


def search(question: str, paper: Optional[str] = None) -> list[dict]:
    """Raw semantic search — top chunks with scores, no LLM."""
    return agent.invoke(_base(action="search", query=question, paper=paper))["results"]


_LEARN_Q = (
    "Give a spoken post-mortem of this research paper: the problem it set out to "
    "solve, the approach and methods used, the key findings and contributions, and "
    "its main limitations. Write flowing prose of about 6 to 8 sentences — no bullet "
    "points or markdown — suitable to be read aloud as an audio summary."
)


def learn(paper: Optional[str] = None) -> str:
    """Spoken post-mortem of `paper` (or all papers), as read-aloud-ready prose."""
    return query(_LEARN_Q, paper=paper)


# ═══════════════════════════════════════════════════════════════════════════════
# CLI
# ═══════════════════════════════════════════════════════════════════════════════

if __name__ == "__main__":
    import sys as _sys

    cmd  = _sys.argv[1] if len(_sys.argv) > 1 else "help"
    args = _sys.argv[2:]

    if cmd == "ingest":
        r = ingest(args)
        p = r["results"][0]["profile"]
        print(f"Ingested:         {r['results'][0]['ingested']}")
        print(f"Total docs:       {r['results'][0]['total_docs']}")
        print(f"Chunk size:       {p['chunk_size']} chars (overlap {p['overlap']})")
        print(f"Pomodoro session: {p['work_minutes']} min")
        print(f"Voice:            #{p['speaker_id']}")

    elif cmd == "query":
        print(query(" ".join(args)))

    elif cmd == "search":
        for h in search(" ".join(args)):
            print(f"[{h['score']:.3f}] ({h['meta'].get('source','?')})  {h['text'][:120]}…")

    elif cmd == "start":
        r = agent.invoke(_base(action="start"))
        t = r["timer"]
        print(f"Timer started — mode: {t['mode']}, duration: {t['settings']['work_duration']//60} min")

    else:
        print(__doc__)
